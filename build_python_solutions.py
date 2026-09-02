from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = 'Python_Control_Flow_Solutions.docx'

def shade(paragraph, fill='F2F4F7'):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); ppr.append(shd)

def font(run, name, size, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:ascii'), name)
    run._element.rPr.rFonts.set(qn('w:hAnsi'), name)
    run.font.size = Pt(size); run.bold = bold
    if color: run.font.color.rgb = RGBColor.from_string(color)

def code(doc, text):
    p = doc.add_paragraph(style='Code Block')
    p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(.15); p.paragraph_format.right_indent = Inches(.15)
    shade(p)
    r = p.add_run(text); font(r, 'Consolas', 8.5)

def add_problem(doc, number, prompt, solution):
    p = doc.add_paragraph(style='Heading 2')
    p.add_run(f'{number}. {prompt}')
    code(doc, solution.strip())

doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Inches(1)
sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(.492)

styles = doc.styles
normal = styles['Normal']; normal.font.name = 'Calibri'; normal._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri'); normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri'); normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.25
for style_name, size, color, before, after in [('Heading 1',16,'2E74B5',18,10), ('Heading 2',13,'2E74B5',14,7), ('Heading 3',12,'1F4D78',10,5)]:
    s=styles[style_name]; s.font.name='Calibri'; s._element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); s._element.rPr.rFonts.set(qn('w:hAnsi'),'Calibri'); s.font.size=Pt(size); s.font.color.rgb=RGBColor.from_string(color); s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after)
if 'Code Block' not in [s.name for s in styles]:
    cb=styles.add_style('Code Block', WD_STYLE_TYPE.PARAGRAPH)
else: cb=styles['Code Block']
cb.font.name='Consolas'; cb._element.rPr.rFonts.set(qn('w:ascii'),'Consolas'); cb._element.rPr.rFonts.set(qn('w:hAnsi'),'Consolas'); cb.font.size=Pt(8.5)

p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(4)
r=p.add_run('PYTHON CONTROL FLOW'); font(r,'Calibri',24,True,'0B2545')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(18)
r=p.add_run('Solved Exercises: If / Elif / Else, Nested Conditions, and Loops'); font(r,'Calibri',12,False,'555555')
p=doc.add_paragraph('Each exercise below includes a complete, corrected Python program. Run each solution separately.'); p.paragraph_format.space_after=Pt(14)

doc.add_heading('1. If-Else Statements', level=1)
ifelse = [
('Check whether a number is even or odd', '''num = int(input("Enter a number: "))

if num % 2 == 0:
    print("The number is even")
else:
    print("The number is odd")'''),
('Check whether student marks are pass or fail', '''marks = int(input("Enter your marks: "))

if marks >= 40:
    print("Pass")
else:
    print("Fail")'''),
('Check voting eligibility', '''age = int(input("Enter your age: "))

if age >= 18:
    print("You are eligible to vote")
else:
    print("You are not eligible to vote")'''),
('Find the greater of two numbers', '''a = int(input("Enter first number: "))
b = int(input("Enter second number: "))

if a > b:
    print("First number is greater")
elif b > a:
    print("Second number is greater")
else:
    print("Both numbers are equal")'''),
('Check whether a number is positive, negative, or zero', '''num = int(input("Enter a number: "))

if num > 0:
    print("The number is positive")
elif num < 0:
    print("The number is negative")
else:
    print("The number is zero")'''),
('ATM: check withdrawal amount', '''balance = 5000
amount = int(input("Enter withdrawal amount: "))

if 0 < amount <= balance:
    balance -= amount
    print("Withdrawal successful")
    print("Remaining balance:", balance)
else:
    print("Invalid amount or insufficient balance")'''),
('Electricity: check free or chargeable usage', '''units = int(input("Enter electricity units consumed: "))

if units <= 100:
    print("No extra charge")
else:
    print("Electricity charges will be applied")'''),
('Online shopping: free delivery', '''amount = float(input("Enter your shopping amount: "))

if amount >= 500:
    print("Congratulations! You get free delivery.")
else:
    print("Delivery charge of Rs. 50 will be added.")''')]
for i,(q,s) in enumerate(ifelse,1): add_problem(doc,i,q,s)

doc.add_heading('2. If-Elif-Else Statements', level=1)
elifprobs=[
('ATM withdrawal with validation', '''balance = 10000
amount = int(input("Enter withdrawal amount: "))

if amount <= 0:
    print("Invalid Amount")
elif amount <= balance:
    balance -= amount
    print("Withdrawal Successful")
    print("Remaining Balance:", balance)
else:
    print("Insufficient Balance")'''),
('Display a student grade from marks', '''marks = int(input("Enter your marks: "))

if marks < 0 or marks > 100:
    print("Invalid marks")
elif marks >= 90:
    print("Grade A")
elif marks >= 75:
    print("Grade B")
elif marks >= 60:
    print("Grade C")
elif marks >= 40:
    print("Grade D")
else:
    print("Fail")'''),
('Calculate an electricity bill by unit slab', '''units = int(input("Enter electricity units consumed: "))

if units < 0:
    print("Invalid unit count")
elif units <= 100:
    bill = units * 3
elif units <= 200:
    bill = units * 5
else:
    bill = units * 7

if units >= 0:
    print("Electricity Bill: Rs.", bill)''')]
for i,(q,s) in enumerate(elifprobs,1): add_problem(doc,i,q,s)

doc.add_heading('3. Nested If-Else Statements', level=1)
nested=[
('Online shopping discount for members', '''member = input("Are you a member? (yes/no): ").strip().lower()
amount = float(input("Enter purchase amount: "))

if member == "yes":
    if amount >= 5000:
        discount = amount * 0.20
    else:
        discount = amount * 0.10
else:
    discount = 0

print("Discount: Rs.", discount)
print("Final Amount: Rs.", amount - discount)'''),
('Unlock a mobile phone using PIN and fingerprint', '''pin = int(input("Enter PIN: "))

if pin == 1234:
    fingerprint = input("Fingerprint verified? (yes/no): ").strip().lower()
    if fingerprint == "yes":
        print("Phone Unlocked")
    else:
        print("Fingerprint Verification Failed")
else:
    print("Incorrect PIN")'''),
('Check online examination eligibility', '''registered = input("Are you registered? (yes/no): ").strip().lower()

if registered == "yes":
    attendance = float(input("Enter attendance percentage: "))
    if attendance >= 75:
        print("Student is eligible for the exam")
    else:
        print("Student is not eligible due to low attendance")
else:
    print("Student is not registered")''')]
for i,(q,s) in enumerate(nested,1): add_problem(doc,i,q,s)

doc.add_heading('4. For Loops', level=1)
forprobs=[
('Homework: print squares from 1 to 10', '''for number in range(1, 11):
    print(number, "squared =", number ** 2)'''),
('Homework: find the factorial of a number', '''n = int(input("Enter a non-negative number: "))
factorial = 1

if n < 0:
    print("Factorial is not defined for negative numbers")
else:
    for number in range(1, n + 1):
        factorial *= number
    print("Factorial =", factorial)'''),
('Use break: stop when number 5 is found', '''for i in range(1, 11):
    if i == 5:
        print("Number 5 found. Stopping the loop.")
        break
    print(i)'''),
('Use continue: skip even numbers', '''for i in range(1, 11):
    if i % 2 == 0:
        continue
    print(i)'''),
('Skip number 3 and stop at number 8', '''for i in range(1, 11):
    if i == 3:
        continue
    if i == 8:
        print("Number 8 found. Loop stopped.")
        break
    print(i)'''),
('Calculate total and average marks for five subjects', '''total = 0

for i in range(1, 6):
    marks = float(input("Enter marks of Subject " + str(i) + ": "))
    total += marks

average = total / 5
print("Total Marks =", total)
print("Average Marks =", average)'''),
('Calculate three days of electricity usage and bill', '''total_units = 0

for i in range(1, 4):
    units = int(input("Enter units consumed on Day " + str(i) + ": "))
    total_units += units

bill = total_units * 8
print("Total Units Consumed =", total_units)
print("Electricity Bill = Rs.", bill)''')]
for i,(q,s) in enumerate(forprobs,1): add_problem(doc,i,q,s)

doc.add_heading('5. While Loops', level=1)
whileprobs=[
('Print numbers from 1 to 10', '''i = 1
while i <= 10:
    print(i)
    i += 1'''),
('Find the sum of the first N natural numbers', '''n = int(input("Enter a positive number: "))
i = 1
total = 0

while i <= n:
    total += i
    i += 1

print("Sum =", total)'''),
('Reverse a given number', '''num = int(input("Enter a non-negative number: "))
original = num
reverse = 0

while num > 0:
    digit = num % 10
    reverse = reverse * 10 + digit
    num //= 10

print("Reversed Number =", reverse)'''),
('Count the digits in a number', '''num = abs(int(input("Enter a number: ")))
count = 1 if num == 0 else 0

while num > 0:
    count += 1
    num //= 10

print("Total Digits =", count)'''),
('Check whether a number is prime', '''num = int(input("Enter a number: "))
i = 2
is_prime = num > 1

while i * i <= num:
    if num % i == 0:
        is_prime = False
        break
    i += 1

if is_prime:
    print("Prime Number")
else:
    print("Not a Prime Number")'''),
('Allow three attempts to enter the correct PIN', '''correct_pin = 1234
attempt = 1

while attempt <= 3:
    pin = int(input("Enter PIN: "))
    if pin == correct_pin:
        print("Login Successful")
        break
    print("Wrong PIN")
    attempt += 1
else:
    print("Account Blocked")'''),
('Password validation until the correct password is entered', '''password = "Python123"

while True:
    user = input("Enter Password: ")
    if user == password:
        print("Access Granted")
        break
    print("Incorrect Password")'''),
('Check whether a number is a palindrome', '''num = int(input("Enter a non-negative number: "))
original = num
reverse = 0

while num > 0:
    digit = num % 10
    reverse = reverse * 10 + digit
    num //= 10

if original == reverse:
    print("Palindrome Number")
else:
    print("Not a Palindrome Number")''')]
for i,(q,s) in enumerate(whileprobs,1): add_problem(doc,i,q,s)

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run('Python Control Flow - Solved Exercises'); font(r,'Calibri',9,False,'666666')
doc.save(OUT)
